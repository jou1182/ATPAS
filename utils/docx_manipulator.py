#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pathlib import Path
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph


# ---------------------------------------------------------------------------
# Open / Save
# ---------------------------------------------------------------------------

def open_docx(path: str | Path) -> Document:
    """Open a .docx file and return a Document object."""
    return Document(str(path))


def save_docx(doc: Document, path: str | Path) -> None:
    """Save a Document, creating parent directories as needed."""
    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))


def new_docx() -> Document:
    """Return a blank Document pre-configured for Arabic RTL content."""
    doc = Document()
    _configure_rtl_document(doc)
    return doc


def _configure_rtl_document(doc: Document) -> None:
    """Set document-level RTL defaults (settings.xml + Normal style).

    Two levels are needed for full compatibility:
    1. ``<w:bidi>`` in ``settings.xml`` — tells Word the document is bidirectional.
    2. ``<w:bidi>`` in the Normal paragraph style's ``<w:pPr>`` — makes every
       paragraph that inherits Normal default to right-to-left.
    Paragraphs with explicit alignment (e.g. CENTER on the cover) are unaffected
    because paragraph-level settings override style defaults.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH as _ALN

    # ── 1. Document settings: mark document as bidirectional ─────────────
    settings_el = doc.settings.element
    if settings_el.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        settings_el.append(bidi)

    # ── 2. Normal style: default direction = RTL, alignment = RIGHT ──────
    try:
        normal = doc.styles["Normal"]
        pPr = normal.element.get_or_add_pPr()
        if pPr.find(qn("w:bidi")) is None:
            b = OxmlElement("w:bidi")
            b.set(qn("w:val"), "1")
            pPr.append(b)
        # Only set if not already defined (don't override explicit overrides)
        if normal.paragraph_format.alignment is None:
            normal.paragraph_format.alignment = _ALN.RIGHT
    except (KeyError, AttributeError):
        pass


# ---------------------------------------------------------------------------
# Trial watermark
# ---------------------------------------------------------------------------

def add_trial_watermark(doc: Document) -> None:
    """
    يضيف إشعار «نسخة تجريبية» في نهاية المستند عند استخدام ترخيص تجريبي (يوم واحد).
    يظهر كفاصل واضح في نهاية الوثيقة حتى لا تُستخدم في مشاريع حقيقية.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    # سطر فاصل
    sep = doc.add_paragraph()
    sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sep = sep.add_run("─" * 55)
    run_sep.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
    run_sep.font.size = Pt(9)

    # نص التحذير
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("⚠  نسخة تجريبية — TRIAL VERSION  ⚠")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)

    # تفاصيل
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_note = note.add_run(
        "هذه الوثيقة صادرة من نسخة تجريبية مجانية من نظام ATPAS.\n"
        "للحصول على نسخة كاملة تواصل مع المطوّر: jou1182@gmail.com"
    )
    run_note.font.size = Pt(10)
    run_note.font.color.rgb = RGBColor(0x88, 0x00, 0x00)
    run_note.font.italic = True

    # سطر فاصل سفلي
    sep2 = doc.add_paragraph()
    sep2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sep2 = sep2.add_run("─" * 55)
    run_sep2.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
    run_sep2.font.size = Pt(9)


# ---------------------------------------------------------------------------
# Paragraph extraction
# ---------------------------------------------------------------------------

def get_paragraphs(doc: Document) -> List[Paragraph]:
    """Return all paragraphs in document body order."""
    return list(doc.paragraphs)


def extract_sections(doc: Document) -> List[Dict]:
    """
    Parse the document into a list of section dicts:
        {"heading": str, "level": int, "paragraphs": [str], "style": str}

    Heading paragraphs (Heading 1/2/3) start a new section.
    Non-heading paragraphs accumulate into the current section.
    """
    sections: List[Dict] = []
    current: Optional[Dict] = None

    for para in doc.paragraphs:
        style_name = para.style.name if para.style else "Normal"
        text = para.text.strip()

        if style_name.startswith("Heading"):
            try:
                level = int(style_name.split()[-1])
            except (ValueError, IndexError):
                level = 1
            current = {"heading": text, "level": level, "paragraphs": [], "style": style_name}
            sections.append(current)
        else:
            if current is None:
                current = {"heading": "", "level": 0, "paragraphs": [], "style": "Normal"}
                sections.append(current)
            if text:
                current["paragraphs"].append(text)

    return sections


# ---------------------------------------------------------------------------
# Style utilities
# ---------------------------------------------------------------------------

def set_paragraph_font(
    para: Paragraph,
    family: str,
    size_pt: float,
    bold: bool = False,
    italic: bool = False,
    color_hex: Optional[str] = None,
) -> None:
    """Apply font properties to every run in a paragraph."""
    for run in para.runs:
        run.font.name = family
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        if color_hex:
            r, g, b = _hex_to_rgb(color_hex)
            run.font.color.rgb = RGBColor(r, g, b)


def set_rtl(para: Paragraph) -> None:
    """Enable right-to-left direction on a paragraph."""
    pPr = para._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
    bidi.set(qn("w:val"), "1")


def clear_document(doc: Document) -> None:
    """Remove all paragraphs and tables from a document body."""
    from docx.oxml.ns import qn as _qn
    body = doc.element.body
    for child in list(body):
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag in ("p", "tbl"):
            body.remove(child)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _hex_to_rgb(hex_color: str) -> Tuple[int, int, int]:
    """Convert '#RRGGBB' to (R, G, B) integers."""
    h = hex_color.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
