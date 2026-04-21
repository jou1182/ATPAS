#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Core document builder: assembles sections from the codes registry into a
formatted Word document with TOC placeholder, page numbers, and owner style.
"""

import logging
import time
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from engine.dependency_resolver import DependencyResolver
from engine.formatter import Formatter
from engine.logger import generate_audit_trail
from engine.style_applier import StyleApplier
from engine.validator import Validator
from utils.content_library import ContentLibrary
from utils.docx_manipulator import new_docx, save_docx
from utils.image_processor import embed_image
from utils.json_manager import load_json

logger = logging.getLogger(__name__)

_METADATA_DIR = Path(".")
_STYLE_DIR = Path("templates/style_templates")
_SOURCE_DOCS_DIR = Path("templates/source_documents")


class Builder:
    """
    Assembles a technical proposal Word document from selected activity codes.

    Usage:
        registry = load_json("codes_registry.json")
        builder = Builder(registry["codes"])
        builder.build(
            selected_codes=["001-SUR-BASE", "002-EXC-FINE", ...],
            project_id="wastewater",
            owner_id="nwc",
            output_path="output/generated_documents/proposal.docx"
        )
    """

    def __init__(
        self,
        codes: Dict[str, Dict],
        metadata_dir: str | Path = _METADATA_DIR,
        style_dir: str | Path = _STYLE_DIR,
        source_docs_dir: str | Path = _SOURCE_DOCS_DIR,
    ):
        self._codes = codes
        self._metadata_dir = Path(metadata_dir)
        self._style_dir = Path(style_dir)
        self._validator = Validator(codes)
        self._resolver = DependencyResolver(codes)
        self._style_applier = StyleApplier(style_dir)
        self._content_lib = ContentLibrary(source_docs_dir)
        self._project_metadata: Dict[str, Dict] = {}

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def build(
        self,
        selected_codes: List[str],
        project_id: str,
        owner_id: str,
        output_path: str | Path,
        skip_validation: bool = False,
    ) -> Tuple[bool, Optional[str]]:
        """
        Build the proposal document.

        Returns:
            (success, error_message_ar)
            error_message_ar is None on success.

        Validation warnings (e.g. missing mandatory codes) are logged but do
        not abort the build.  Hard errors (forbidden codes, inactive codes)
        cause build failure.
        """
        output_path = Path(output_path)
        start = time.monotonic()

        # --- Validate ---
        if not skip_validation:
            is_valid, errors, warnings = self._validator.validate(
                selected_codes, owner_id, project_id
            )
            for w in warnings:
                logger.warning("Validation warning: %s", w)
            if not is_valid:
                msg = "فشل التحقق من الأكواد:\n" + "\n".join(f"• {e}" for e in errors)
                logger.error("Build aborted — validation failed:\n%s", "\n".join(errors))
                self._write_audit(
                    selected_codes, project_id, owner_id, output_path,
                    "failure", msg, 0, 0
                )
                return False, msg

        # --- Resolve & order codes ---
        ordered_codes = self._resolver.resolve(selected_codes)
        project_meta = self._load_project_metadata(project_id)

        # --- Assemble document ---
        doc = new_docx()
        style_spec = self._load_style_spec(owner_id)
        formatter = Formatter(style_spec) if style_spec else None

        self._add_cover(doc, project_id, owner_id, ordered_codes, formatter)
        self._add_toc_placeholder(doc)
        self._add_sections(doc, ordered_codes, project_meta, formatter)
        self._style_applier.apply_style(doc, owner_id)

        # --- Save ---
        save_docx(doc, output_path)

        elapsed = time.monotonic() - start
        file_size = output_path.stat().st_size if output_path.exists() else 0

        logger.info(
            "Build complete: %s  (%.2fs, %d bytes)",
            output_path, elapsed, file_size
        )
        self._write_audit(
            selected_codes, project_id, owner_id, output_path,
            "success", None, elapsed, file_size
        )
        return True, None

    def estimate_pages(self, selected_codes: List[str]) -> int:
        """Estimate total page count for the selected codes."""
        return sum(
            self._codes[c].get("page_count", 0)
            for c in selected_codes
            if c in self._codes
        )

    def estimate_images(self, selected_codes: List[str]) -> int:
        """Estimate total image count for the selected codes."""
        return sum(
            self._codes[c].get("image_count", 0)
            for c in selected_codes
            if c in self._codes
        )

    def content_summary(self, selected_codes: List[str]) -> Dict[str, str]:
        """
        Return a dict mapping each code to its content status:
            "library"     → real .docx file found in content library
            "placeholder" → no source file yet (will use metadata text)
        """
        return {
            code_id: ("library" if self._content_lib.exists(code_id) else "placeholder")
            for code_id in selected_codes
            if code_id in self._codes
        }

    # ------------------------------------------------------------------
    # Document assembly
    # ------------------------------------------------------------------

    def _add_cover(
        self,
        doc: Document,
        project_id: str,
        owner_id: str,
        ordered_codes: List[str],
        formatter: Optional[Formatter],
    ) -> None:
        from datetime import date
        project_meta = self._load_project_metadata(project_id)
        project_name = project_meta.get("project_metadata", {}).get(
            "name_ar", project_id
        )

        owner_names = {"nwc": "الشركة الوطنية للمياه", "makkah": "أمانة العاصمة المقدسة", "moh": "وزارة الإسكان"}
        owner_name = owner_names.get(owner_id, owner_id)

        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_para.add_run("عرض فني")
        run.font.size = Pt(24)
        run.bold = True

        subtitle_para = doc.add_paragraph()
        subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_para.add_run(project_name).font.size = Pt(16)

        owner_para = doc.add_paragraph()
        owner_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        owner_para.add_run(f"مُقدَّم إلى: {owner_name}").font.size = Pt(14)

        date_para = doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_para.add_run(f"التاريخ: {date.today().strftime('%Y-%m-%d')}").font.size = Pt(12)

        pages = self.estimate_pages(ordered_codes)
        stats_para = doc.add_paragraph()
        stats_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        stats_para.add_run(
            f"عدد الأكواد: {len(ordered_codes)}  |  الصفحات التقديرية: {pages}"
        ).font.size = Pt(11)

        doc.add_page_break()

    def _add_toc_placeholder(self, doc: Document) -> None:
        """Insert a TOC field that Word regenerates on open."""
        toc_heading = doc.add_heading("فهرس المحتويات", level=1)
        toc_heading.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Insert TOC field via raw OOXML
        para = doc.add_paragraph()
        run = para.add_run()
        fldChar_begin = OxmlElement("w:fldChar")
        fldChar_begin.set(qn("w:fldCharType"), "begin")
        fldChar_begin.set(qn("w:dirty"), "true")

        instrText = OxmlElement("w:instrText")
        instrText.set(qn("xml:space"), "preserve")
        instrText.text = ' TOC \\o "1-3" \\h \\z \\u '

        fldChar_end = OxmlElement("w:fldChar")
        fldChar_end.set(qn("w:fldCharType"), "end")

        run._r.append(fldChar_begin)
        run._r.append(instrText)
        run._r.append(fldChar_end)

        doc.add_page_break()

    def _add_sections(
        self,
        doc: Document,
        ordered_codes: List[str],
        project_meta: Dict,
        formatter: Optional[Formatter],
    ) -> None:
        """
        Add one section per code in sequence order.

        Priority:
          1. Real content from content library (templates/source_documents/{code_id}.docx)
          2. Metadata-based placeholder (activity name + description)
        """
        activities = {
            a["code_id"]: a
            for a in project_meta.get("activity_sequence", [])
        }

        for code_id in ordered_codes:
            code = self._codes.get(code_id)
            if not code:
                continue

            activity = activities.get(code_id, {})
            name_ar = code.get("activity_name_ar", code_id)
            name_en = code.get("activity_name_en", "")
            phase = activity.get("phase_name_ar", "")

            # Section heading — always shown regardless of content source
            heading_text = f"{code_id}  —  {name_ar}"
            if formatter:
                formatter.add_heading(doc, heading_text, level=1)
            else:
                doc.add_heading(heading_text, level=1)

            # Try content library first
            has_real_content = self._content_lib.insert_into(doc, code_id)

            if not has_real_content:
                # Fallback: placeholder from metadata
                if name_en:
                    if formatter:
                        formatter.add_heading(doc, name_en, level=3)
                    else:
                        doc.add_heading(name_en, level=3)

                description = activity.get("description", "")
                if description:
                    if formatter:
                        formatter.add_body_paragraph(doc, description)
                    else:
                        doc.add_paragraph(description)

                pages = code.get("page_count", 0)
                images = code.get("image_count", 0)
                meta_para = doc.add_paragraph(
                    f"الصفحات: {pages}  |  الصور: {images}  |  المرحلة: {phase}"
                )
                meta_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

                if images > 0:
                    img_ph = doc.add_paragraph(
                        f"[{images} صورة — ضع ملف {code_id}.docx في templates/source_documents/]"
                    )
                    img_ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                logger.info("Used content library for %s", code_id)

            doc.add_paragraph()  # section spacer

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    def _load_project_metadata(self, project_id: str) -> Dict:
        if project_id in self._project_metadata:
            return self._project_metadata[project_id]
        path = self._metadata_dir / f"{project_id}_project_metadata.json"
        try:
            meta = load_json(path)
        except FileNotFoundError:
            meta = {"activity_sequence": []}
        self._project_metadata[project_id] = meta
        return meta

    def _load_style_spec(self, owner_id: str) -> Optional[Dict]:
        path = self._style_dir / f"{owner_id}_style.json"
        try:
            return load_json(path)
        except FileNotFoundError:
            return None

    def _write_audit(
        self,
        codes: List[str],
        project_id: str,
        owner_id: str,
        output_path: Path,
        status: str,
        error: Optional[str],
        elapsed: float,
        file_size: int,
    ) -> None:
        try:
            generate_audit_trail({
                "selected_codes": codes,
                "project_id": project_id,
                "owner_id": owner_id,
                "output_path": str(output_path),
                "status": status,
                "error": error,
                "processing_time_seconds": round(elapsed, 3),
                "file_size_bytes": file_size,
            })
        except Exception as exc:
            logger.warning("Audit trail write failed: %s", exc)
