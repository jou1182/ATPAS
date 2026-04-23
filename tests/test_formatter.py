#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""Unit tests for engine/formatter.py.

Tests cover: style spec parsing, heading/body paragraph creation, RTL direction,
spacing, font application, and fallback behaviour on missing keys.

Run from the project root:
    python -m pytest tests/test_formatter.py -v
"""

from __future__ import annotations

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from engine.formatter import Formatter


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_spec(
    *,
    family: str = "Arial",
    size: int = 12,
    bold: bool = False,
    italic: bool = False,
    heading_color: str | None = None,
    text_color: str | None = None,
    rtl: bool = True,
) -> dict:
    """Build a minimal style spec dict for Formatter."""
    spec: dict = {
        "fonts": {
            "body": {"family": family, "size": size, "bold": bold, "italic": italic},
            "heading1": {"family": family, "size": size + 6, "bold": True},
            "heading2": {"family": family, "size": size + 3, "bold": True},
            "heading3": {"family": family, "size": size + 1, "bold": False},
        },
        "rtl_direction": rtl,
    }
    colors: dict = {}
    if heading_color:
        colors["heading"] = heading_color
    if text_color:
        colors["text"] = text_color
    if colors:
        spec["colors"] = colors
    return spec


# ---------------------------------------------------------------------------
# Initialisation
# ---------------------------------------------------------------------------

class TestFormatterInit:
    def test_empty_spec_does_not_crash(self) -> None:
        f = Formatter({})
        assert f._rtl is True     # default

    def test_reads_rtl_false(self) -> None:
        f = Formatter({"rtl_direction": False})
        assert f._rtl is False

    def test_reads_rtl_true(self) -> None:
        f = Formatter({"rtl_direction": True})
        assert f._rtl is True

    def test_missing_fonts_key_returns_empty_dict(self) -> None:
        f = Formatter({})
        assert f._fonts == {}

    def test_missing_colors_key_returns_empty_dict(self) -> None:
        f = Formatter({})
        assert f._colors == {}


# ---------------------------------------------------------------------------
# add_heading
# ---------------------------------------------------------------------------

class TestAddHeading:
    def test_returns_paragraph(self) -> None:
        doc = Document()
        f = Formatter(_make_spec())
        para = f.add_heading(doc, "عنوان اختبار", level=1)
        assert para is not None
        assert "عنوان اختبار" in para.text

    def test_heading_text_in_document(self) -> None:
        doc = Document()
        f = Formatter(_make_spec())
        f.add_heading(doc, "Test Heading", level=1)
        texts = [p.text for p in doc.paragraphs]
        assert "Test Heading" in texts

    def test_rtl_true_sets_right_alignment(self) -> None:
        doc = Document()
        f = Formatter(_make_spec(rtl=True))
        para = f.add_heading(doc, "RTL heading", level=1)
        assert para.alignment == WD_ALIGN_PARAGRAPH.RIGHT

    def test_rtl_false_does_not_force_right(self) -> None:
        doc = Document()
        f = Formatter(_make_spec(rtl=False))
        para = f.add_heading(doc, "LTR heading", level=1)
        # When RTL is off, Formatter must not override alignment to RIGHT
        assert para.alignment != WD_ALIGN_PARAGRAPH.RIGHT

    def test_level_2_heading(self) -> None:
        doc = Document()
        f = Formatter(_make_spec())
        para = f.add_heading(doc, "Level 2", level=2)
        assert "Level 2" in para.text

    def test_level_3_heading(self) -> None:
        doc = Document()
        f = Formatter(_make_spec())
        para = f.add_heading(doc, "Level 3", level=3)
        assert "Level 3" in para.text

    def test_color_override_accepted(self) -> None:
        """Passing color_hex must not raise."""
        doc = Document()
        f = Formatter(_make_spec())
        # Should not raise even if color is supplied
        para = f.add_heading(doc, "Colored", level=1, color_hex="#FF0000")
        assert para is not None


# ---------------------------------------------------------------------------
# add_body_paragraph
# ---------------------------------------------------------------------------

class TestAddBodyParagraph:
    def test_returns_paragraph(self) -> None:
        doc = Document()
        f = Formatter(_make_spec())
        para = f.add_body_paragraph(doc, "نص الجسم")
        assert "نص الجسم" in para.text

    def test_text_present_in_document(self) -> None:
        doc = Document()
        f = Formatter(_make_spec())
        f.add_body_paragraph(doc, "Body text")
        texts = [p.text for p in doc.paragraphs]
        assert "Body text" in texts

    def test_rtl_true_sets_right_alignment(self) -> None:
        doc = Document()
        f = Formatter(_make_spec(rtl=True))
        para = f.add_body_paragraph(doc, "RTL body")
        assert para.alignment == WD_ALIGN_PARAGRAPH.RIGHT

    def test_rtl_false_does_not_force_right(self) -> None:
        doc = Document()
        f = Formatter(_make_spec(rtl=False))
        para = f.add_body_paragraph(doc, "LTR body")
        assert para.alignment != WD_ALIGN_PARAGRAPH.RIGHT

    def test_multiple_paragraphs_independent(self) -> None:
        doc = Document()
        f = Formatter(_make_spec())
        f.add_body_paragraph(doc, "First")
        f.add_body_paragraph(doc, "Second")
        texts = [p.text for p in doc.paragraphs]
        assert "First" in texts
        assert "Second" in texts


# ---------------------------------------------------------------------------
# set_paragraph_spacing
# ---------------------------------------------------------------------------

class TestSetParagraphSpacing:
    def test_space_before_after(self) -> None:
        doc = Document()
        f = Formatter(_make_spec())
        para = doc.add_paragraph("spacing test")
        f.set_paragraph_spacing(para, space_before_pt=10, space_after_pt=8)
        assert para.paragraph_format.space_before == Pt(10)
        assert para.paragraph_format.space_after == Pt(8)

    def test_line_spacing_applied_when_given(self) -> None:
        doc = Document()
        f = Formatter(_make_spec())
        para = doc.add_paragraph("line spacing test")
        f.set_paragraph_spacing(para, line_spacing_pt=14)
        assert para.paragraph_format.line_spacing == Pt(14)

    def test_no_line_spacing_when_not_given(self) -> None:
        doc = Document()
        f = Formatter(_make_spec())
        para = doc.add_paragraph("no line spacing")
        # Default line_spacing_pt is None → should not set line_spacing
        f.set_paragraph_spacing(para, space_before_pt=6, space_after_pt=6)
        # paragraph_format.line_spacing may be None or AUTO — should not be Pt(x)
        ls = para.paragraph_format.line_spacing
        assert ls is None or not isinstance(ls, int) or ls <= 0


# ---------------------------------------------------------------------------
# format_document (batch format all paragraphs)
# ---------------------------------------------------------------------------

class TestFormatDocument:
    def test_does_not_raise_on_empty_doc(self) -> None:
        doc = Document()
        f = Formatter(_make_spec())
        f.format_document(doc)   # must not raise

    def test_does_not_raise_on_headings_and_body(self) -> None:
        doc = Document()
        doc.add_heading("Heading 1", level=1)
        doc.add_heading("Heading 2", level=2)
        doc.add_paragraph("Body paragraph")
        f = Formatter(_make_spec())
        f.format_document(doc)   # must not raise

    def test_arabic_text_survives_format(self) -> None:
        doc = Document()
        doc.add_paragraph("أعمال مواسير الصرف الصحي")
        f = Formatter(_make_spec(rtl=True))
        f.format_document(doc)
        texts = [p.text for p in doc.paragraphs]
        assert "أعمال مواسير الصرف الصحي" in texts


# ---------------------------------------------------------------------------
# format_paragraph
# ---------------------------------------------------------------------------

class TestFormatParagraph:
    def test_formats_body_paragraph(self) -> None:
        doc = Document()
        para = doc.add_paragraph("test")
        f = Formatter(_make_spec())
        f.format_paragraph(para, style_key="body")   # must not raise

    def test_rtl_sets_right_alignment(self) -> None:
        doc = Document()
        para = doc.add_paragraph("RTL")
        f = Formatter(_make_spec(rtl=True))
        f.format_paragraph(para, style_key="body")
        assert para.alignment == WD_ALIGN_PARAGRAPH.RIGHT


# ---------------------------------------------------------------------------
# Font fallback behaviour
# ---------------------------------------------------------------------------

class TestFontFallback:
    def test_missing_style_key_falls_back_to_body(self) -> None:
        """Formatter should not raise when a requested style_key is missing."""
        spec = {
            "fonts": {"body": {"family": "Arial", "size": 12}},
            "rtl_direction": True,
        }
        doc = Document()
        f = Formatter(spec)
        para = f.add_heading(doc, "No heading spec", level=1)
        assert para is not None   # fell back gracefully

    def test_completely_empty_fonts_no_crash(self) -> None:
        """Formatter with empty fonts dict must not raise."""
        doc = Document()
        f = Formatter({"fonts": {}, "rtl_direction": False})
        para = f.add_body_paragraph(doc, "Empty fonts")
        assert "Empty fonts" in para.text
