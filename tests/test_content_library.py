#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""TDD tests for utils/content_library.py — ContentLibrary.

Covers: find(), exists(), register(), unregister(), list_available(),
        invalidate_cache(), and the fuzzy-match consistency contract.

Run from the project root:
    python -m pytest tests/test_content_library.py -v
"""

import json
from pathlib import Path

import pytest
from docx import Document

from utils.content_library import ContentLibrary


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_minimal_docx(path: Path, text: str = "محتوى تجريبي") -> None:
    """Write a minimal valid .docx to path."""
    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))


def _make_lib(tmp_path: Path, registry: dict | None = None) -> ContentLibrary:
    """Create a ContentLibrary pointing at a tmp source_docs dir."""
    src = tmp_path / "source_documents"
    src.mkdir()
    reg_path = tmp_path / "content_registry.json"
    if registry is not None:
        reg_path.write_text(json.dumps(registry, ensure_ascii=False), encoding="utf-8")
    return ContentLibrary(source_docs_dir=src, registry_path=reg_path)


# ---------------------------------------------------------------------------
# find() — basic lookup
# ---------------------------------------------------------------------------

class TestFind:
    def test_returns_none_for_empty_dir(self, tmp_path):
        lib = _make_lib(tmp_path)
        assert lib.find("001-SUR-BASE") is None

    def test_exact_filename_match(self, tmp_path):
        lib = _make_lib(tmp_path)
        src = tmp_path / "source_documents"
        _make_minimal_docx(src / "001-SUR-BASE.docx")

        result = lib.find("001-SUR-BASE")

        assert result is not None
        assert result.name == "001-SUR-BASE.docx"

    def test_fuzzy_hyphen_insensitive_match(self, tmp_path):
        """File named '001surbase.docx' must be found for code '001-SUR-BASE'."""
        lib = _make_lib(tmp_path)
        src = tmp_path / "source_documents"
        _make_minimal_docx(src / "001surbase.docx")

        result = lib.find("001-SUR-BASE")

        assert result is not None, "find() should match via fuzzy/normalized lookup"

    def test_case_insensitive_match(self, tmp_path):
        lib = _make_lib(tmp_path)
        src = tmp_path / "source_documents"
        _make_minimal_docx(src / "001-sur-base.docx")

        result = lib.find("001-SUR-BASE")

        assert result is not None

    def test_registry_entry_takes_priority(self, tmp_path):
        """An explicit registry path beats a same-named file in source_docs."""
        src = tmp_path / "source_documents"
        src.mkdir()
        _make_minimal_docx(src / "001-SUR-BASE.docx", text="from file")

        custom_dir = tmp_path / "custom"
        custom_dir.mkdir()
        custom_file = custom_dir / "custom.docx"
        _make_minimal_docx(custom_file, text="from registry")

        registry = {"001-SUR-BASE": str(custom_file)}
        reg_path = tmp_path / "content_registry.json"
        reg_path.write_text(json.dumps(registry), encoding="utf-8")
        lib = ContentLibrary(source_docs_dir=src, registry_path=reg_path)

        result = lib.find("001-SUR-BASE")
        assert result == custom_file

    def test_registry_missing_file_falls_through_to_scan(self, tmp_path):
        """If registry points to a nonexistent file, fall through to filesystem scan."""
        src = tmp_path / "source_documents"
        src.mkdir()
        _make_minimal_docx(src / "001-SUR-BASE.docx")

        registry = {"001-SUR-BASE": str(tmp_path / "missing.docx")}
        reg_path = tmp_path / "content_registry.json"
        reg_path.write_text(json.dumps(registry), encoding="utf-8")
        lib = ContentLibrary(source_docs_dir=src, registry_path=reg_path)

        result = lib.find("001-SUR-BASE")
        assert result is not None  # falls through to exact file match
        assert result.name == "001-SUR-BASE.docx"

    def test_returns_none_when_no_match(self, tmp_path):
        lib = _make_lib(tmp_path)
        src = tmp_path / "source_documents"
        _make_minimal_docx(src / "OTHER-CODE.docx")

        assert lib.find("001-SUR-BASE") is None


# ---------------------------------------------------------------------------
# exists() — consistency with find()
# ---------------------------------------------------------------------------

class TestExists:
    def test_returns_false_for_empty_dir(self, tmp_path):
        lib = _make_lib(tmp_path)
        assert lib.exists("001-SUR-BASE") is False

    def test_exact_filename_true(self, tmp_path):
        lib = _make_lib(tmp_path)
        src = tmp_path / "source_documents"
        _make_minimal_docx(src / "001-SUR-BASE.docx")
        assert lib.exists("001-SUR-BASE") is True

    def test_fuzzy_match_consistent_with_find(self, tmp_path):
        """exists() must return True whenever find() returns a path.

        This test exposes the fuzzy-match consistency bug: file is named
        '001surbase.docx' but code queried as '001-SUR-BASE'.  find() uses
        fuzzy scan and returns the file; exists() must agree.
        """
        lib = _make_lib(tmp_path)
        src = tmp_path / "source_documents"
        _make_minimal_docx(src / "001surbase.docx")

        found = lib.find("001-SUR-BASE")
        assert found is not None, "Precondition: find() must locate the fuzzy file"

        # The key assertion — exists() and find() must agree
        assert lib.exists("001-SUR-BASE") is True, (
            "exists() returned False even though find() found the file. "
            "exists() does not handle fuzzy-normalized filenames."
        )

    def test_unknown_code_false(self, tmp_path):
        lib = _make_lib(tmp_path)
        src = tmp_path / "source_documents"
        _make_minimal_docx(src / "OTHER.docx")
        assert lib.exists("001-SUR-BASE") is False

    def test_registry_registered_code_true(self, tmp_path):
        """A code in the registry (with existing file) must be True."""
        src = tmp_path / "source_documents"
        src.mkdir()
        docx_file = tmp_path / "manual.docx"
        _make_minimal_docx(docx_file)

        registry = {"MY-CODE": str(docx_file)}
        reg_path = tmp_path / "content_registry.json"
        reg_path.write_text(json.dumps(registry), encoding="utf-8")
        lib = ContentLibrary(source_docs_dir=src, registry_path=reg_path)

        assert lib.exists("MY-CODE") is True

    def test_registry_missing_file_false(self, tmp_path):
        """Registry entry pointing to nonexistent file must return False."""
        src = tmp_path / "source_documents"
        src.mkdir()

        registry = {"MY-CODE": str(tmp_path / "ghost.docx")}
        reg_path = tmp_path / "content_registry.json"
        reg_path.write_text(json.dumps(registry), encoding="utf-8")
        lib = ContentLibrary(source_docs_dir=src, registry_path=reg_path)

        assert lib.exists("MY-CODE") is False


# ---------------------------------------------------------------------------
# register() + unregister() — runtime mutations
# ---------------------------------------------------------------------------

class TestRegisterUnregister:
    def test_register_makes_exists_true(self, tmp_path):
        lib = _make_lib(tmp_path)
        docx_file = tmp_path / "manual.docx"
        _make_minimal_docx(docx_file)

        assert lib.exists("NEW-CODE") is False
        lib.register("NEW-CODE", docx_file, save=False)
        assert lib.exists("NEW-CODE") is True

    def test_register_makes_find_work(self, tmp_path):
        lib = _make_lib(tmp_path)
        docx_file = tmp_path / "manual.docx"
        _make_minimal_docx(docx_file)

        lib.register("NEW-CODE", docx_file, save=False)
        assert lib.find("NEW-CODE") == docx_file

    def test_register_persists_to_registry_json(self, tmp_path):
        lib = _make_lib(tmp_path)
        docx_file = tmp_path / "manual.docx"
        _make_minimal_docx(docx_file)
        reg_path = tmp_path / "content_registry.json"

        lib.register("PERSIST-CODE", docx_file, save=True)

        saved = json.loads(reg_path.read_text(encoding="utf-8"))
        assert "PERSIST-CODE" in saved

    def test_unregister_makes_exists_false(self, tmp_path):
        lib = _make_lib(tmp_path)
        docx_file = tmp_path / "manual.docx"
        _make_minimal_docx(docx_file)
        lib.register("CODE-X", docx_file, save=False)
        assert lib.exists("CODE-X") is True

        lib.unregister("CODE-X", save=False)
        assert lib.exists("CODE-X") is False

    def test_unregister_nonexistent_code_is_noop(self, tmp_path):
        lib = _make_lib(tmp_path)
        lib.unregister("DOES-NOT-EXIST", save=False)  # must not raise

    def test_cache_invalidated_after_register(self, tmp_path):
        """exists() must reflect changes immediately after register()."""
        lib = _make_lib(tmp_path)
        docx_file = tmp_path / "late.docx"
        _make_minimal_docx(docx_file)

        # Prime the cache with a miss
        assert lib.exists("LATE-CODE") is False

        # Register and verify cache is invalidated
        lib.register("LATE-CODE", docx_file, save=False)
        assert lib.exists("LATE-CODE") is True


# ---------------------------------------------------------------------------
# invalidate_cache()
# ---------------------------------------------------------------------------

class TestInvalidateCache:
    def test_invalidate_causes_rescan(self, tmp_path):
        """After invalidation, a newly added file should be detected."""
        lib = _make_lib(tmp_path)
        src = tmp_path / "source_documents"

        # Prime cache — no files
        assert lib.exists("005-HND-FIN") is False

        # Add file on disk without going through library API
        _make_minimal_docx(src / "005-HND-FIN.docx")

        # Without invalidation, cache might still say False
        lib.invalidate_cache()

        # Now it must say True
        assert lib.exists("005-HND-FIN") is True


# ---------------------------------------------------------------------------
# list_available()
# ---------------------------------------------------------------------------

class TestListAvailable:
    def test_empty_dir_returns_empty_list(self, tmp_path):
        lib = _make_lib(tmp_path)
        assert lib.list_available() == []

    def test_lists_file_stems(self, tmp_path):
        lib = _make_lib(tmp_path)
        src = tmp_path / "source_documents"
        _make_minimal_docx(src / "001-SUR-BASE.docx")
        _make_minimal_docx(src / "002-EXC-FINE.docx")

        result = lib.list_available()
        assert "001-SUR-BASE" in result
        assert "002-EXC-FINE" in result

    def test_returns_sorted(self, tmp_path):
        lib = _make_lib(tmp_path)
        src = tmp_path / "source_documents"
        _make_minimal_docx(src / "Z-CODE.docx")
        _make_minimal_docx(src / "A-CODE.docx")
        _make_minimal_docx(src / "M-CODE.docx")

        result = lib.list_available()
        assert result == sorted(result)

    def test_includes_registry_codes(self, tmp_path):
        docx_file = tmp_path / "external.docx"
        _make_minimal_docx(docx_file)
        lib = _make_lib(tmp_path, registry={"REG-CODE": str(docx_file)})

        result = lib.list_available()
        assert "REG-CODE" in result


# ---------------------------------------------------------------------------
# insert_into()
# ---------------------------------------------------------------------------

class TestInsertInto:
    def test_returns_false_when_no_source(self, tmp_path):
        lib = _make_lib(tmp_path)
        doc = Document()
        result = lib.insert_into(doc, "MISSING-CODE")
        assert result is False

    def test_returns_true_and_copies_content(self, tmp_path):
        lib = _make_lib(tmp_path)
        src = tmp_path / "source_documents"
        _make_minimal_docx(src / "001-SUR-BASE.docx", text="النص الأصلي")

        target = Document()
        result = lib.insert_into(target, "001-SUR-BASE")

        assert result is True
        all_text = " ".join(p.text for p in target.paragraphs)
        assert "النص الأصلي" in all_text
