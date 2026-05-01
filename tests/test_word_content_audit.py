#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pathlib import Path

from docx import Document

from utils.content_library import ContentLibrary
from utils.docx_manipulator import set_rtl
from utils.word_content_audit import (
    audit_code_content,
    audit_docx_file,
    audit_selected_content,
    flatten_issues,
)


def _save_docx(path: Path, paragraphs: list[str], *, rtl: bool = True) -> None:
    doc = Document()
    for text in paragraphs:
        para = doc.add_paragraph(text)
        if rtl:
            set_rtl(para)
    doc.save(str(path))


def test_audit_missing_content_is_warning(tmp_path: Path) -> None:
    src = tmp_path / "source_documents"
    src.mkdir()
    lib = ContentLibrary(source_docs_dir=src, registry_path=tmp_path / "content_registry.json")

    audit = audit_code_content("001-SUR-BASE", {"activity_name_ar": "مسح عام"}, lib)

    assert audit.exists is False
    assert audit.errors == []
    assert audit.warnings
    assert "لا يوجد ملف Word" in audit.warnings[0].message_ar


def test_audit_valid_docx_is_ok(tmp_path: Path) -> None:
    path = tmp_path / "001-SUR-BASE.docx"
    _save_docx(path, ["يشمل هذا البند أعمال المسح والتحقق من النقاط المرجعية."])

    audit = audit_docx_file(
        path,
        code_id="001-SUR-BASE",
        code_data={"activity_name_ar": "مسح وتثبيت نقاط عامة"},
    )

    assert audit.exists is True
    assert audit.ok is True
    assert audit.paragraph_count == 1


def test_audit_flags_duplicate_heading(tmp_path: Path) -> None:
    path = tmp_path / "001-SUR-BASE.docx"
    _save_docx(path, ["مسح وتثبيت نقاط عامة", "يشمل هذا البند أعمال المسح."])

    audit = audit_docx_file(
        path,
        code_id="001-SUR-BASE",
        code_data={"activity_name_ar": "مسح وتثبيت نقاط عامة"},
    )

    messages = "\n".join(issue.message_ar for issue in audit.warnings)
    assert "عنوان النشاط" in messages


def test_audit_does_not_flag_opening_paragraph_that_mentions_activity(tmp_path: Path) -> None:
    path = tmp_path / "001-APP-DES.docx"
    _save_docx(path, [
        "تعتبر مرحلة إعداد واعتماد المخططات التفصيلية هي المركز الأساسي للأعمال الهندسية، "
        "حيث يتم فيها تحويل المخططات التصميمية إلى وثائق تنفيذية تفصيلية قابلة للتطبيق في الموقع."
    ])

    audit = audit_docx_file(
        path,
        code_id="001-APP-DES",
        code_data={"activity_name_ar": "اعتماد المخططات التفصيلية"},
    )

    messages = "\n".join(issue.message_ar for issue in audit.warnings)
    assert "عنوان النشاط" not in messages


def test_audit_flags_empty_docx_as_error(tmp_path: Path) -> None:
    path = tmp_path / "001-SUR-BASE.docx"
    Document().save(str(path))

    audit = audit_docx_file(path, code_id="001-SUR-BASE", code_data={})

    assert audit.errors
    assert "فارغ" in audit.errors[0].message_ar


def test_audit_selected_skips_custom_codes(tmp_path: Path) -> None:
    src = tmp_path / "source_documents"
    src.mkdir()
    lib = ContentLibrary(source_docs_dir=src, registry_path=tmp_path / "content_registry.json")
    codes = {
        "999-CUS-001": {"is_custom": True, "activity_name_ar": "بند مخصص"},
        "001-SUR-BASE": {"activity_name_ar": "مسح عام"},
    }

    audits = audit_selected_content(codes, ["999-CUS-001", "001-SUR-BASE"], lib)
    issues = flatten_issues(audits)

    assert [audit.code_id for audit in audits] == ["001-SUR-BASE"]
    assert len(issues) == 1
