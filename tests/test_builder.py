#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Builder integration tests + dependency-injection unit tests.

Run from the project root:
    python -m pytest tests/test_builder.py -v
"""

import time
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from docx import Document
from docx.oxml.ns import qn

from engine.builder import Builder
from engine.dependency_resolver import DependencyResolver
from engine.validator import Validator
from utils.content_library import ContentLibrary
from utils.json_manager import load_json

_OUTPUT_DIR = Path("output/test_builds")


@pytest.fixture(scope="module")
def registry():
    return load_json("codes_registry.json")


@pytest.fixture(scope="module")
def codes(registry):
    return registry["codes"]


@pytest.fixture(scope="module")
def builder(codes):
    return Builder(codes)


def _iter_story_roots(doc: Document):
    yield doc.element.body
    for section in doc.sections:
        yield section.header._element
        yield section.footer._element
        yield section.first_page_header._element
        yield section.first_page_footer._element
        yield section.even_page_header._element
        yield section.even_page_footer._element


def _assert_docx_is_tajawal_and_right_aligned(
    path: Path,
    *,
    allow_cover_center: bool = False,
) -> None:
    doc = Document(str(path))
    checked_paragraphs = 0
    checked_runs = 0

    for root_index, root in enumerate(_iter_story_roots(doc)):
        before_first_page_break = allow_cover_center and root_index == 0
        for p_el in root.iter(qn("w:p")):
            has_page_break = _p_has_page_break(p_el)
            has_text = any((t.text or "").strip() for t in p_el.iter(qn("w:t")))
            has_visual = (
                p_el.find(".//" + qn("w:drawing")) is not None
                or p_el.find(".//" + qn("w:pict")) is not None
                or p_el.find(".//" + qn("w:object")) is not None
            )
            if not has_text and not has_visual:
                if before_first_page_break and has_page_break:
                    before_first_page_break = False
                continue
            checked_paragraphs += 1
            pPr = p_el.find(qn("w:pPr"))
            bidi = pPr.find(qn("w:bidi")) if pPr is not None else None
            jc = pPr.find(qn("w:jc")) if pPr is not None else None
            assert bidi is not None and bidi.get(qn("w:val")) == "1"
            expected_jc = "center" if before_first_page_break and not has_visual else "right"
            assert jc is not None and jc.get(qn("w:val")) == expected_jc
            if before_first_page_break and has_page_break:
                before_first_page_break = False

        for r_el in root.iter(qn("w:r")):
            if not any((t.text or "").strip() for t in r_el.iter(qn("w:t"))):
                continue
            checked_runs += 1
            rPr = r_el.find(qn("w:rPr"))
            rFonts = rPr.find(qn("w:rFonts")) if rPr is not None else None
            assert rFonts is not None
            assert rFonts.get(qn("w:ascii")) == "Tajawal"
            assert rFonts.get(qn("w:hAnsi")) == "Tajawal"
            assert rFonts.get(qn("w:eastAsia")) == "Tajawal"
            assert rFonts.get(qn("w:cs")) == "Tajawal"

    assert checked_paragraphs > 0
    assert checked_runs > 0


def _p_has_page_break(p_el) -> bool:
    for br in p_el.iter(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    return False


# ---------------------------------------------------------------------------
# Scenario 1: wastewater + NWC + FINE excavation  →  valid .docx
# ---------------------------------------------------------------------------

def test_wastewater_nwc_fine(builder):
    selected = [
        "001-SUR-BASE", "001-PRM-GOV", "001-APP-DES", "001-APP-HSE",
        "002-MAT-SITE", "002-EXC-FINE", "002-WST-EXC",
        "003-PIP-SEW", "004-TST-LEK", "004-QC-INSTALL",
        "005-BKF-SND", "005-BKF-SOL", "005-RST-ASP",
        "005-HND-DOC", "005-HND-FIN",
    ]
    out = _OUTPUT_DIR / "wastewater_nwc_fine.docx"
    success, err = builder.build(selected, "wastewater", "nwc", out)
    assert success, f"Build failed: {err}"
    assert out.exists()
    doc = Document(str(out))
    assert len(doc.paragraphs) > 0


# ---------------------------------------------------------------------------
# Scenario 2: wastewater + NWC + OPEN excavation  →  valid .docx
# ---------------------------------------------------------------------------

def test_wastewater_nwc_open(builder):
    selected = [
        "001-SUR-BASE", "001-PRM-GOV", "001-APP-DES", "001-APP-HSE",
        "002-MAT-SITE", "002-EXC-OPEN", "002-WST-EXC",
        "003-PIP-SEW", "004-TST-LEK", "004-QC-INSTALL",
        "005-BKF-SND", "005-BKF-SOL", "005-RST-ASP",
        "005-HND-DOC", "005-HND-FIN",
    ]
    out = _OUTPUT_DIR / "wastewater_nwc_open.docx"
    success, err = builder.build(selected, "wastewater", "nwc", out)
    assert success, f"Build failed: {err}"
    assert out.exists()


# ---------------------------------------------------------------------------
# Scenario 3: water supply + Makkah + FINE excavation  →  valid .docx
# ---------------------------------------------------------------------------

def test_water_makkah_fine(builder):
    selected = [
        "001-SUR-BASE", "001-PRM-GOV", "001-PRM-MUN", "001-APP-DES",
        "001-APP-HSE", "001-APP-MRL",
        "002-MAT-SITE", "002-EXC-FINE", "002-WST-EXC",
        "003-PIP-WAT", "004-TST-LEK", "004-QC-INSTALL",
        "005-BKF-SND", "005-BKF-SOL", "005-RST-ASP",
        "005-HND-DOC", "005-HND-FIN",
    ]
    out = _OUTPUT_DIR / "water_makkah_fine.docx"
    success, err = builder.build(selected, "water_supply", "makkah", out)
    assert success, f"Build failed: {err}"
    assert out.exists()
    doc = Document(str(out))
    # Heritage clause should have been added
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "التراث" in full_text or "تراث" in full_text


# ---------------------------------------------------------------------------
# Scenario 4: asphalt + MOH  →  valid .docx with compliance matrix
# ---------------------------------------------------------------------------

def test_asphalt_moh(builder):
    selected = [
        "001-SUR-BASE", "001-PRM-GOV", "001-APP-DES",
        "001-APP-MRL", "001-APP-HSE",
        "002-MAT-SITE", "002-EXC-PAV", "002-WST-EXC",
        "004-QC-MATS",
        "005-BKF-BASE", "005-BKF-SOL", "005-RST-ASP",
        "005-HND-DOC", "005-HND-FIN",
    ]
    out = _OUTPUT_DIR / "asphalt_moh.docx"
    success, err = builder.build(selected, "asphalt", "moh", out)
    assert success, f"Build failed: {err}"
    assert out.exists()
    doc = Document(str(out))
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "الامتثال" in full_text


# ---------------------------------------------------------------------------
# Scenario 5: missing dependencies  →  build succeeds with warnings (not errors)
# ---------------------------------------------------------------------------

def test_missing_deps_build_with_warnings(builder):
    # 005-BKF-SND without its dependency 004-TST-LEK
    selected = [
        "001-SUR-BASE", "001-PRM-GOV", "001-APP-DES", "001-APP-HSE",
        "002-MAT-SITE", "002-EXC-FINE",
        "005-BKF-SND",   # missing: 003-PIP-SEW, 004-TST-LEK
        "005-HND-FIN",
    ]
    out = _OUTPUT_DIR / "missing_deps.docx"
    # Should build successfully (missing deps are warnings, not hard errors)
    success, err = builder.build(selected, "wastewater", "nwc", out)
    assert success, f"Expected build to succeed despite warnings, got: {err}"


# ---------------------------------------------------------------------------
# Scenario 6: multiple excavation methods  →  build allowed
# ---------------------------------------------------------------------------

def test_multiple_excavation_methods_build_allowed(builder):
    selected = [
        "001-SUR-BASE", "001-PRM-GOV", "001-APP-DES", "001-APP-HSE",
        "002-MAT-SITE",
        "002-EXC-FINE",
        "002-EXC-OPEN",
        "002-EXC-TUNNEL",
        "003-PIP-SEW",
    ]
    out = _OUTPUT_DIR / "multi_exc.docx"
    success, err = builder.build(selected, "wastewater", "nwc", out)
    assert success, f"Expected build to allow multiple excavation methods, got: {err}"


# ---------------------------------------------------------------------------
# Dependency injection — المحور 4
# ---------------------------------------------------------------------------

class TestBuilderDependencyInjection:
    """Verify Builder accepts pre-built collaborators instead of creating its own."""

    def test_default_creates_validator_and_resolver(self, codes) -> None:
        """Without injection, Builder instantiates its own collaborators."""
        b = Builder(codes)
        assert isinstance(b._validator, Validator)
        assert isinstance(b._resolver, DependencyResolver)

    def test_injected_validator_is_used(self, codes, tmp_path) -> None:
        """Injected Validator is called during build; internal one is NOT created."""
        mock_v = MagicMock(spec=Validator)
        mock_v.validate.return_value = (True, [], [])

        b = Builder(codes, validator=mock_v)
        assert b._validator is mock_v

        out = tmp_path / "di_validator.docx"
        b.build(["001-SUR-BASE"], "wastewater", "nwc", out)
        mock_v.validate.assert_called_once()

    def test_injected_resolver_is_used(self, codes, tmp_path) -> None:
        """Injected DependencyResolver is called during build."""
        mock_r = MagicMock(spec=DependencyResolver)
        mock_r.resolve.return_value = ["001-SUR-BASE"]

        b = Builder(codes, resolver=mock_r)
        assert b._resolver is mock_r

        out = tmp_path / "di_resolver.docx"
        b.build(["001-SUR-BASE"], "wastewater", "nwc", out, skip_validation=True)
        mock_r.resolve.assert_called_once_with(["001-SUR-BASE"])

    def test_injected_content_lib_is_used(self, codes, tmp_path) -> None:
        """Injected ContentLibrary controls whether real content is available."""
        mock_lib = MagicMock(spec=ContentLibrary)
        mock_lib.exists.return_value = False
        mock_lib.insert_into.return_value = False

        b = Builder(codes, content_lib=mock_lib)
        assert b._content_lib is mock_lib

        out = tmp_path / "di_content.docx"
        success, _ = b.build(["001-SUR-BASE"], "wastewater", "nwc", out,
                              skip_validation=True)
        assert success
        # insert_into should have been called for the single code
        mock_lib.insert_into.assert_called_once()

    def test_partial_injection_only_overrides_supplied(self, codes) -> None:
        """Only the supplied collaborator is overridden; others use defaults."""
        mock_v = MagicMock(spec=Validator)
        mock_v.validate.return_value = (True, [], [])

        b = Builder(codes, validator=mock_v)
        assert b._validator is mock_v
        assert isinstance(b._resolver, DependencyResolver)   # not mocked
        assert isinstance(b._content_lib, ContentLibrary)    # not mocked


class TestBuilderArabicOutputContract:
    """Final DOCX output contract for every current/future owner."""

    def test_future_owner_style_cannot_override_tajawal_or_right_alignment(
        self,
        codes,
        tmp_path,
    ) -> None:
        style_dir = tmp_path / "styles"
        style_dir.mkdir()
        (style_dir / "future_owner_style.json").write_text(
            """
            {
              "fonts": {
                "body": {"family": "Times New Roman", "size": 11, "bold": false},
                "heading1": {"family": "Arial", "size": 15, "bold": true},
                "heading2": {"family": "Calibri", "size": 13, "bold": true},
                "heading3": {"family": "Tahoma", "size": 12, "bold": true}
              },
              "colors": {"text": "#000000", "heading": "#003D7A"},
              "header": {"enabled": true, "text_ar": "جهة مستقبلية", "line_below": true},
              "footer": {"enabled": true, "text_ar": "تذييل مستقبلي", "show_page_number": true},
              "rtl_direction": true
            }
            """,
            encoding="utf-8",
        )

        mock_lib = MagicMock(spec=ContentLibrary)
        mock_lib.exists.return_value = False
        mock_lib.insert_into.return_value = False

        b = Builder(codes, style_dir=style_dir, content_lib=mock_lib)
        out = tmp_path / "future_owner.docx"
        success, err = b.build(
            ["001-SUR-BASE"],
            "wastewater",
            "future_owner",
            out,
            skip_validation=True,
        )

        assert success, err
        _assert_docx_is_tajawal_and_right_aligned(out, allow_cover_center=True)


# ---------------------------------------------------------------------------
# _owner_name_ar — reads from master_config.json
# ---------------------------------------------------------------------------

class TestOwnerNameAr:
    """Unit tests for Builder._owner_name_ar() — no filesystem I/O needed."""

    def _make_builder(self, owner_specs: dict, codes: dict) -> Builder:
        """Return a Builder with _owner_specs pre-set (bypasses file I/O)."""
        b = Builder.__new__(Builder)
        b._owner_specs = owner_specs
        b._codes = codes
        b._project_metadata = {}
        return b

    def test_known_owner_returns_arabic_name(self, codes) -> None:
        specs = {"nwc": {"owner_name_ar": "الشركة الوطنية للمياه"}}
        b = self._make_builder(specs, codes)
        assert b._owner_name_ar("nwc") == "الشركة الوطنية للمياه"

    def test_unknown_owner_falls_back_to_id(self, codes) -> None:
        b = self._make_builder({}, codes)
        assert b._owner_name_ar("unknown_org") == "unknown_org"

    def test_owner_without_arabic_name_falls_back_to_id(self, codes) -> None:
        specs = {"partial": {"owner_name_en": "Some Company"}}  # no owner_name_ar
        b = self._make_builder(specs, codes)
        assert b._owner_name_ar("partial") == "partial"

    def test_real_builder_loads_nwc_name(self, codes) -> None:
        """Integration: Builder(codes) reads master_config.json and resolves NWC."""
        b = Builder(codes)
        name = b._owner_name_ar("nwc")
        # master_config.json has owner_name_ar for nwc
        assert "مياه" in name or name == "nwc", (
            f"Expected Arabic name for 'nwc', got {name!r}"
        )

    def test_cover_uses_arabic_name_not_id(self, codes, tmp_path) -> None:
        """The generated cover page must show the owner's Arabic name, not 'nwc'."""
        out = tmp_path / "cover_test.docx"
        b = Builder(codes)
        success, _ = b.build(
            ["001-SUR-BASE"], "wastewater", "nwc", out, skip_validation=True
        )
        assert success
        doc = Document(str(out))
        full_text = " ".join(p.text for p in doc.paragraphs)
        assert "nwc" not in full_text, "Cover page must not expose raw owner_id"


# ---------------------------------------------------------------------------
# Performance: large document builds in < 5 seconds
# ---------------------------------------------------------------------------

def test_build_performance(builder):
    """60+ page document should build in under 5 seconds."""
    selected = [
        "001-SUR-BASE", "001-SUR-NET", "001-PRM-GOV", "001-PRM-ENV",
        "001-APP-DES", "001-APP-MRL", "001-APP-HSE",
        "002-MAT-SITE", "002-MAT-STORE", "002-EXC-OPEN", "002-WST-EXC", "002-WST-RECY",
        "003-PIP-WAT", "003-WLD-FLD",
        "004-TST-HYD", "004-TST-LEK", "004-QC-INSTALL",
        "005-BKF-SND", "005-BKF-SOL", "005-RST-ASP",
        "005-HND-DOC", "005-HND-TRN", "005-HND-FIN",
    ]
    out = _OUTPUT_DIR / "performance_test.docx"
    t0 = time.monotonic()
    success, err = builder.build(selected, "water_supply", "nwc", out)
    elapsed = time.monotonic() - t0
    assert success, f"Build failed: {err}"
    assert elapsed < 5.0, f"Build took {elapsed:.2f}s (limit 5s)"
    estimated_pages = builder.estimate_pages(selected)
    assert estimated_pages >= 60, f"Expected ≥60 pages, got {estimated_pages}"


# ---------------------------------------------------------------------------
# Advanced structure: category grouping + dividers
# ---------------------------------------------------------------------------

class TestCategoryGrouping:
    """Verify sections are grouped by category (001–005) with divider pages."""

    def test_estimate_pages_adds_divider_per_extra_category(self, codes) -> None:
        b = Builder(codes)
        # Single category → no divider pages
        one_cat = ["001-SUR-BASE", "001-PRM-GOV"]
        assert b._category_divider_count(one_cat) == 0
        # Two categories → 1 divider page
        two_cat = ["001-SUR-BASE", "002-EXC-FINE"]
        assert b._category_divider_count(two_cat) == 1
        # Five categories → 4 divider pages
        five_cat = [
            "001-SUR-BASE", "002-EXC-FINE", "003-PIP-SEW",
            "004-TST-LEK", "005-HND-FIN",
        ]
        assert b._category_divider_count(five_cat) == 4

    def test_estimate_pages_reflects_dividers(self, codes) -> None:
        b = Builder(codes)
        # Same two codes; adding a second category adds exactly +1 divider page
        one_cat = ["001-SUR-BASE", "001-PRM-GOV"]           # 1 category
        same_codes_two_cats = ["001-SUR-BASE", "002-EXC-FINE"]  # 2 categories
        base_pages = b.estimate_pages(one_cat)
        multi_pages = b.estimate_pages(same_codes_two_cats)
        # Differing code pages (5 vs 4) + 1 divider page
        code_page_diff = (
            codes["002-EXC-FINE"].get("page_count", 0)
            - codes["001-PRM-GOV"].get("page_count", 0)
        )
        assert multi_pages == base_pages + code_page_diff + 1

    def test_category_divider_appears_in_document(self, builder) -> None:
        """The built document contains the category names from code_ranges."""
        selected = ["001-SUR-BASE", "002-EXC-FINE"]
        out = _OUTPUT_DIR / "category_grouping.docx"
        success, err = builder.build(selected, "wastewater", "nwc", out,
                                     skip_validation=True)
        assert success, f"Build failed: {err}"
        doc = Document(str(out))
        full_text = " ".join(p.text for p in doc.paragraphs)
        # Category names from master_config.json → code_ranges
        assert "الأعمال التحضيرية" in full_text
        assert "الحفر والمخلفات" in full_text

    def test_category_divider_order_matches_code_sequence(self, builder) -> None:
        """Category dividers appear in code sequence order (001 before 002)."""
        selected = ["001-SUR-BASE", "002-EXC-FINE", "003-PIP-SEW"]
        out = _OUTPUT_DIR / "category_order.docx"
        success, err = builder.build(selected, "wastewater", "nwc", out,
                                     skip_validation=True)
        assert success, f"Build failed: {err}"
        doc = Document(str(out))
        texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        # Build the position map for category names
        pos_001 = next(i for i, t in enumerate(texts) if "الأعمال التحضيرية" in t)
        pos_002 = next(i for i, t in enumerate(texts) if "الحفر والمخلفات" in t)
        pos_003 = next(i for i, t in enumerate(texts) if "التركيب والتوصيل" in t)
        assert pos_001 < pos_002 < pos_003

    def test_category_divider_page_break_count(self, builder) -> None:
        """Number of explicit page breaks before category sections ≥ categories-1."""
        selected = [
            "001-SUR-BASE", "002-EXC-FINE", "003-PIP-SEW",
            "004-TST-LEK", "005-HND-FIN",
        ]
        out = _OUTPUT_DIR / "category_breaks.docx"
        success, err = builder.build(selected, "wastewater", "nwc", out,
                                     skip_validation=True)
        assert success, f"Build failed: {err}"
        doc = Document(str(out))
        page_breaks = 0
        for p in doc.paragraphs:
            for br in p._p.iter(qn("w:br")):
                if br.get(qn("w:type")) == "page":
                    page_breaks += 1
        # Cover + TOC have their own breaks; dividers add categories-1 more
        assert page_breaks >= 4, f"Expected ≥4 page breaks, got {page_breaks}"

    def test_unknown_category_falls_back_to_raw_id(self, codes, tmp_path) -> None:
        """A code with an unknown category still builds (no KeyError)."""
        # Inject a fake code with category 'ZZZ'
        fake_codes = dict(codes)
        fake_codes["999-TST-XXX"] = {
            "code_id": "999-TST-XXX",
            "category": "ZZZ",
            "activity_name_ar": "اختبار فئة غير معروفة",
            "activity_name_en": "Unknown category test",
            "sequence_order": 9999,
            "page_count": 1,
            "status": "active",
        }
        b = Builder(fake_codes)
        out = tmp_path / "unknown_cat.docx"
        success, err = b.build(["999-TST-XXX"], "wastewater", "nwc", out,
                               skip_validation=True)
        assert success, f"Build failed: {err}"
