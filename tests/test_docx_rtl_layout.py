#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from utils.docx_manipulator import enforce_document_rtl


def _ppr(para):
    return para._p.get_or_add_pPr()


def _bidi_value(para) -> str | None:
    bidi = _ppr(para).find(qn("w:bidi"))
    return bidi.get(qn("w:val")) if bidi is not None else None


def _jc_value(para) -> str | None:
    jc = _ppr(para).find(qn("w:jc"))
    return jc.get(qn("w:val")) if jc is not None else None


def test_enforce_document_rtl_patches_body_table_header_and_footer() -> None:
    doc = Document()
    body_para = doc.add_paragraph("فقرة عربية في جسم المستند")
    body_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    table = doc.add_table(rows=1, cols=1)
    cell_para = table.cell(0, 0).paragraphs[0]
    cell_para.text = "فقرة عربية داخل جدول"
    cell_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    header_para = doc.sections[0].header.paragraphs[0]
    header_para.text = "رأس الصفحة"
    header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer_para = doc.sections[0].footer.paragraphs[0]
    footer_para.text = "تذييل الصفحة"
    footer_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    enforce_document_rtl(doc)

    for para in (body_para, cell_para, header_para, footer_para):
        assert _bidi_value(para) == "1"
        assert _jc_value(para) == "right"


def test_enforce_document_rtl_preserves_cover_center_before_first_page_break() -> None:
    doc = Document()
    centered = doc.add_paragraph("عنوان غلاف في المنتصف")
    centered.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    enforce_document_rtl(doc)

    assert _bidi_value(centered) == "1"
    assert _jc_value(centered) == "center"


def test_enforce_document_rtl_moves_center_after_cover_to_right() -> None:
    doc = Document()
    cover = doc.add_paragraph("غلاف")
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    body = doc.add_paragraph("نص بعد الغلاف")
    body.alignment = WD_ALIGN_PARAGRAPH.CENTER

    enforce_document_rtl(doc)

    assert _jc_value(cover) == "center"
    assert _bidi_value(body) == "1"
    assert _jc_value(body) == "right"


def test_enforce_document_rtl_moves_visual_paragraphs_to_right() -> None:
    doc = Document()
    visual = doc.add_paragraph()
    visual.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = visual.add_run()
    drawing = OxmlElement("w:drawing")
    run._r.append(drawing)

    enforce_document_rtl(doc)

    assert _bidi_value(visual) == "1"
    assert _jc_value(visual) == "right"
