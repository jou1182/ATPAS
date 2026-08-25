#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""Code management dialog for ATPAS.

This dialog lets non-technical users add, edit, activate, and deactivate codes
without opening codes_registry.json directly.
"""

from __future__ import annotations

import copy
import os
import re
from datetime import date
from pathlib import Path
from typing import Any

from PyQt5.QtCore import Qt, pyqtSignal
from PyQt5.QtWidgets import (
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QFormLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QSplitter,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

from ui import theme
from utils.json_manager import save_json
from utils.registry_validator import validate_registry


_REGISTRY_PATH = Path("codes_registry.json")
_SOURCE_DOCS_DIR = Path("templates/source_documents")
_CODE_ID_RE = re.compile(r"^\d{3}-[A-Z]{2,8}-[A-Z]{2,8}$")


class _MultiPickerDialog(QDialog):
    """Small RTL multi-select dialog for projects, owners, networks, or deps."""

    def __init__(
        self,
        title: str,
        options: list[tuple[str, str]],
        selected: list[str],
        parent=None,
    ) -> None:
        super().__init__(parent)
        self.setWindowTitle(title)
        self.setLayoutDirection(Qt.RightToLeft)
        self.setMinimumSize(420, 480)

        root = QVBoxLayout(self)
        root.setContentsMargins(12, 12, 12, 12)
        root.setSpacing(8)

        self._search = QLineEdit()
        self._search.setPlaceholderText("بحث...")
        self._search.setLayoutDirection(Qt.RightToLeft)
        self._search.textChanged.connect(self._filter)
        root.addWidget(self._search)

        self._list = QListWidget()
        self._list.setLayoutDirection(Qt.RightToLeft)
        self._items: list[QListWidgetItem] = []
        selected_set = set(selected)
        for item_id, label in options:
            item = QListWidgetItem(f"{label}  |  {item_id}")
            item.setData(Qt.UserRole, item_id)
            item.setFlags(item.flags() | Qt.ItemIsUserCheckable)
            item.setCheckState(Qt.Checked if item_id in selected_set else Qt.Unchecked)
            self._list.addItem(item)
            self._items.append(item)
        root.addWidget(self._list, stretch=1)

        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        root.addWidget(buttons)

    def selected_ids(self) -> list[str]:
        return [
            item.data(Qt.UserRole)
            for item in self._items
            if item.checkState() == Qt.Checked
        ]

    def _filter(self, text: str) -> None:
        needle = text.strip().lower()
        for item in self._items:
            item.setHidden(needle not in item.text().lower())


class CodeManagerDialog(QDialog):
    """Manage code metadata in codes_registry.json through a guided UI."""

    data_changed = pyqtSignal()

    def __init__(self, registry_data: dict[str, Any], config_data: dict[str, Any], parent=None) -> None:
        super().__init__(parent)
        self.setWindowTitle("إدارة الأكواد")
        self.setLayoutDirection(Qt.RightToLeft)
        self.setMinimumSize(980, 680)

        self._registry_data = registry_data
        self._config_data = config_data
        self._codes: dict[str, dict[str, Any]] = self._registry_data.setdefault("codes", {})
        self._current_code_id: str | None = None
        self._is_new = False
        self._selected_projects: list[str] = []
        self._selected_owners: list[str] = []
        self._selected_networks: list[str] = []
        self._selected_deps: list[str] = []

        self._build_ui()
        self._refresh_list()

    def _build_ui(self) -> None:
        root = QVBoxLayout(self)
        root.setContentsMargins(14, 12, 14, 12)
        root.setSpacing(10)

        title = QLabel("إدارة الأكواد الفنية")
        title.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        title.setStyleSheet(
            f"font-size: 18px; font-weight: 900; color: {theme.HEADER}; "
            f"padding-bottom: 7px; border-bottom: 2px solid {theme.ACCENT};"
        )
        root.addWidget(title)

        hint = QLabel(
            "أضف أو عدّل كودًا من هنا بدون فتح ملف JSON. استخدم التعطيل بدل الحذف للحفاظ على السجل."
        )
        hint.setWordWrap(True)
        hint.setStyleSheet(f"color: {theme.TEXT2}; font-size: 12px;")
        root.addWidget(hint)

        splitter = QSplitter(Qt.Horizontal)
        splitter.setLayoutDirection(Qt.RightToLeft)
        root.addWidget(splitter, stretch=1)

        list_panel = QWidget()
        list_layout = QVBoxLayout(list_panel)
        list_layout.setContentsMargins(0, 0, 0, 0)
        list_layout.setSpacing(8)

        self._search = QLineEdit()
        self._search.setPlaceholderText("بحث بالكود أو الاسم...")
        self._search.setLayoutDirection(Qt.RightToLeft)
        self._search.textChanged.connect(self._refresh_list)
        list_layout.addWidget(self._search)

        self._code_list = QListWidget()
        self._code_list.setLayoutDirection(Qt.RightToLeft)
        self._code_list.currentItemChanged.connect(self._on_selected_item_changed)
        list_layout.addWidget(self._code_list, stretch=1)

        new_btn = QPushButton("إضافة كود جديد")
        new_btn.clicked.connect(self._new_code)
        list_layout.addWidget(new_btn)

        form_panel = QWidget()
        form_layout = QVBoxLayout(form_panel)
        form_layout.setContentsMargins(10, 0, 0, 0)
        form_layout.setSpacing(10)

        form = QFormLayout()
        form.setLabelAlignment(Qt.AlignRight)
        form.setFormAlignment(Qt.AlignRight | Qt.AlignTop)
        form.setSpacing(8)

        self._code_id = QLineEdit()
        self._code_id.setLayoutDirection(Qt.LeftToRight)
        self._code_id.setAlignment(Qt.AlignLeft)
        self._code_id.setPlaceholderText("مثال: 003-PIP-MAIN")
        form.addRow("معرف الكود", self._code_id)

        self._name_ar = QLineEdit()
        self._name_ar.setLayoutDirection(Qt.RightToLeft)
        self._name_ar.setAlignment(Qt.AlignRight)
        form.addRow("الاسم بالعربية", self._name_ar)

        self._name_en = QLineEdit()
        self._name_en.setLayoutDirection(Qt.LeftToRight)
        self._name_en.setAlignment(Qt.AlignLeft)
        form.addRow("الاسم بالإنجليزية", self._name_en)

        self._category = QComboBox()
        for cid, label in [
            ("001", "001 - الأعمال التحضيرية"),
            ("002", "002 - الحفر والمخلفات"),
            ("003", "003 - التركيب والتوصيل"),
            ("004", "004 - الاختبارات والفحوصات"),
            ("005", "005 - الإنهاء والتسليم"),
        ]:
            self._category.addItem(label, cid)
        form.addRow("الفئة", self._category)

        self._phase = QLineEdit()
        self._phase.setLayoutDirection(Qt.LeftToRight)
        self._phase.setAlignment(Qt.AlignLeft)
        self._phase.setPlaceholderText("PIP")
        form.addRow("المرحلة", self._phase)

        self._variation = QLineEdit()
        self._variation.setLayoutDirection(Qt.LeftToRight)
        self._variation.setAlignment(Qt.AlignLeft)
        self._variation.setPlaceholderText("MAIN")
        form.addRow("النوع", self._variation)

        self._sequence = QSpinBox()
        self._sequence.setRange(0, 99999)
        form.addRow("ترتيب الظهور", self._sequence)

        self._pages = QSpinBox()
        self._pages.setRange(0, 999)
        self._pages.setSuffix(" صفحة")
        form.addRow("عدد الصفحات", self._pages)

        self._status = QComboBox()
        self._status.addItem("نشط", "active")
        self._status.addItem("غير نشط", "inactive")
        self._status.addItem("مسودة", "draft")
        self._status.addItem("مؤرشف", "archived")
        form.addRow("الحالة", self._status)

        self._source_document = QLineEdit()
        self._source_document.setLayoutDirection(Qt.LeftToRight)
        self._source_document.setAlignment(Qt.AlignLeft)
        self._source_document.setPlaceholderText("اختياري: 003-PIP-MAIN.docx")
        form.addRow("ملف Word", self._source_document)

        form_layout.addLayout(form)

        picker_row = QHBoxLayout()
        self._projects_btn = QPushButton("المشاريع (0)")
        self._owners_btn = QPushButton("الجهات (0)")
        self._networks_btn = QPushButton("الشبكات (0)")
        self._deps_btn = QPushButton("التبعيات (0)")
        self._projects_btn.clicked.connect(self._pick_projects)
        self._owners_btn.clicked.connect(self._pick_owners)
        self._networks_btn.clicked.connect(self._pick_networks)
        self._deps_btn.clicked.connect(self._pick_deps)
        picker_row.addWidget(self._projects_btn)
        picker_row.addWidget(self._owners_btn)
        picker_row.addWidget(self._networks_btn)
        picker_row.addWidget(self._deps_btn)
        form_layout.addLayout(picker_row)

        self._notes = QTextEdit()
        self._notes.setPlaceholderText("ملاحظات اختيارية...")
        self._notes.setMinimumHeight(90)
        form_layout.addWidget(self._notes)

        action_row = QHBoxLayout()
        self._toggle_btn = QPushButton("تعطيل الكود")
        self._toggle_btn.clicked.connect(self._toggle_status)
        self._save_btn = QPushButton("حفظ التعديل")
        self._save_btn.setObjectName("codeManagerSaveBtn")
        self._save_btn.clicked.connect(self._save_current)
        action_row.addWidget(self._toggle_btn)
        action_row.addStretch()
        action_row.addWidget(self._save_btn)
        form_layout.addLayout(action_row)

        word_row = QHBoxLayout()
        self._word_btn = QPushButton("📄  فتح / إنشاء محتوى Word")
        self._word_btn.setObjectName("codeManagerWordBtn")
        self._word_btn.setEnabled(False)
        self._word_btn.setToolTip(
            "يفتح ملف Word الخاص بهذا الكود في Microsoft Word.\n"
            "إن لم يكن الملف موجوداً، يُنشئ قالباً جاهزاً للتحرير."
        )
        self._word_btn.clicked.connect(self._open_or_create_docx)
        word_row.addStretch()
        word_row.addWidget(self._word_btn)
        form_layout.addLayout(word_row)

        splitter.addWidget(form_panel)
        splitter.addWidget(list_panel)
        splitter.setSizes([650, 330])

        close_row = QHBoxLayout()
        close_row.addStretch()
        close_btn = QPushButton("إغلاق")
        close_btn.clicked.connect(self.accept)
        close_row.addWidget(close_btn)
        root.addLayout(close_row)

        self.setStyleSheet(
            f"""
            QDialog {{
                background: {theme.SURFACE};
                color: {theme.TEXT};
            }}
            QLineEdit, QTextEdit, QComboBox, QListWidget, QSpinBox {{
                background: white;
                border: 1px solid {theme.BORDER2};
                border-radius: 7px;
                padding: 6px;
                color: {theme.TEXT};
            }}
            QPushButton {{
                min-height: 32px;
                border-radius: 7px;
                border: 1px solid {theme.BORDER2};
                background: {theme.SURFACE};
                color: {theme.TEXT};
                font-weight: 700;
                padding: 6px 14px;
            }}
            QPushButton:hover {{
                background: {theme.ACCENT_PALE};
                border-color: {theme.ACCENT};
                color: {theme.ACCENT_DARK};
            }}
            QPushButton#codeManagerSaveBtn {{
                background: {theme.SUCCESS};
                color: white;
                border: none;
                font-weight: 900;
            }}
            QPushButton#codeManagerSaveBtn:hover {{
                background: #236040;
                color: white;
            }}
            QPushButton#codeManagerWordBtn {{
                background: #1F3A56;
                color: white;
                border: none;
                font-weight: 700;
                min-width: 220px;
            }}
            QPushButton#codeManagerWordBtn:hover {{
                background: {theme.HEADER};
                color: #F5D48B;
            }}
            QPushButton#codeManagerWordBtn:disabled {{
                background: {theme.BORDER2};
                color: #888;
            }}
            """
        )

    def _refresh_list(self) -> None:
        query = self._search.text().strip().lower() if hasattr(self, "_search") else ""
        current = self._current_code_id
        self._code_list.blockSignals(True)
        self._code_list.clear()
        for code_id in sorted(self._codes, key=lambda c: self._codes[c].get("sequence_order", 99999)):
            entry = self._codes[code_id]
            name = entry.get("activity_name_ar", "")
            if query and query not in code_id.lower() and query not in str(name).lower():
                continue
            status = entry.get("status", "active")
            status_label = "نشط" if status == "active" else "غير نشط"
            item = QListWidgetItem(f"{code_id}  |  {name}  |  {status_label}")
            item.setData(Qt.UserRole, code_id)
            if status != "active":
                item.setForeground(Qt.gray)
            self._code_list.addItem(item)
            if code_id == current:
                self._code_list.setCurrentItem(item)
        self._code_list.blockSignals(False)
        if self._code_list.currentItem() is None and self._code_list.count():
            self._code_list.setCurrentRow(0)
        elif self._code_list.count() == 0:
            self._clear_form()

    def _on_selected_item_changed(self, current: QListWidgetItem | None, _previous) -> None:
        if current is None:
            return
        self._load_code(current.data(Qt.UserRole))

    def _load_code(self, code_id: str) -> None:
        entry = copy.deepcopy(self._codes.get(code_id, {}))
        self._is_new = False
        self._current_code_id = code_id
        self._word_btn.setEnabled(True)
        self._code_id.setReadOnly(True)
        self._code_id.setText(code_id)
        self._name_ar.setText(str(entry.get("activity_name_ar", "")))
        self._name_en.setText(str(entry.get("activity_name_en", "")))
        self._set_combo_data(self._category, str(entry.get("category", code_id[:3])))
        self._phase.setText(str(entry.get("phase", "")))
        self._variation.setText(str(entry.get("variation", "")))
        self._sequence.setValue(_safe_int(entry.get("sequence_order"), 9999))
        self._pages.setValue(_safe_int(entry.get("page_count"), 0))
        self._set_combo_data(self._status, str(entry.get("status", "active")))
        self._source_document.setText(str(entry.get("source_document", "")))
        self._notes.setPlainText(str(entry.get("notes", "")))
        self._selected_projects = _string_list(entry.get("project_ids"))
        self._selected_owners = _string_list(entry.get("applicable_owners"))
        self._selected_networks = _string_list(entry.get("network_types"))
        self._selected_deps = _string_list(entry.get("dependencies"))
        self._sync_picker_labels()
        self._sync_toggle_label()

    def _new_code(self) -> None:
        self._is_new = True
        self._current_code_id = None
        self._word_btn.setEnabled(False)
        self._code_list.clearSelection()
        self._code_id.setReadOnly(False)
        self._code_id.setText("")
        self._name_ar.setText("")
        self._name_en.setText("")
        self._set_combo_data(self._category, "001")
        self._phase.setText("")
        self._variation.setText("")
        self._sequence.setValue(9999)
        self._pages.setValue(0)
        self._set_combo_data(self._status, "active")
        self._source_document.setText("")
        self._notes.setPlainText("")
        self._selected_projects = []
        self._selected_owners = []
        self._selected_networks = []
        self._selected_deps = []
        self._sync_picker_labels()
        self._sync_toggle_label()

    def _clear_form(self) -> None:
        self._current_code_id = None
        self._word_btn.setEnabled(False)
        self._code_id.setText("")
        self._name_ar.setText("")
        self._name_en.setText("")
        self._notes.setPlainText("")

    def _pick_projects(self) -> None:
        self._selected_projects = self._pick("اختر المشاريع", self._project_options(), self._selected_projects)
        self._sync_picker_labels()

    def _pick_owners(self) -> None:
        self._selected_owners = self._pick("اختر الجهات المالكة", self._owner_options(), self._selected_owners)
        self._sync_picker_labels()

    def _pick_networks(self) -> None:
        self._selected_networks = self._pick("اختر أنواع الشبكات", self._network_options(), self._selected_networks)
        self._sync_picker_labels()

    def _pick_deps(self) -> None:
        current_id = self._code_id.text().strip()
        options = [
            (cid, f"{entry.get('activity_name_ar', '')}")
            for cid, entry in sorted(self._codes.items())
            if cid != current_id
        ]
        self._selected_deps = self._pick("اختر التبعيات", options, self._selected_deps)
        self._sync_picker_labels()

    def _pick(self, title: str, options: list[tuple[str, str]], selected: list[str]) -> list[str]:
        dialog = _MultiPickerDialog(title, options, selected, self)
        if dialog.exec_() == QDialog.Accepted:
            return dialog.selected_ids()
        return selected

    def _save_current(self) -> None:
        code_id = self._code_id.text().strip().upper()
        error = self._validate_form(code_id)
        if error:
            QMessageBox.warning(self, "لا يمكن الحفظ", error)
            return

        old_id = self._current_code_id
        entry = self._entry_from_form(code_id)
        candidate = copy.deepcopy(self._registry_data)
        candidate_codes = candidate.setdefault("codes", {})
        if old_id and old_id != code_id:
            candidate_codes.pop(old_id, None)
        candidate_codes[code_id] = entry

        issues = validate_registry(candidate)
        if issues:
            scope, field, message = issues[0]
            QMessageBox.warning(
                self,
                "فشل التحقق",
                f"يوجد خطأ في بيانات الكود قبل الحفظ:\n[{scope}].{field}: {message}",
            )
            return

        self._registry_data.clear()
        self._registry_data.update(candidate)
        self._codes = self._registry_data["codes"]
        save_json(self._registry_data, _REGISTRY_PATH)
        self._current_code_id = code_id
        self._is_new = False
        self._code_id.setReadOnly(True)
        self._refresh_list()
        self._select_code(code_id)
        self.data_changed.emit()
        QMessageBox.information(self, "تم الحفظ", f"تم حفظ الكود {code_id} بنجاح.")

    def _open_or_create_docx(self) -> None:
        """Open the source .docx for the current code in Word.

        If the file does not exist, create a structured Arabic template first,
        notify the user, then open it.  The file is always opened via
        os.startfile() which delegates to the system's default .docx handler.
        """
        code_id = self._current_code_id
        if not code_id:
            return

        _SOURCE_DOCS_DIR.mkdir(parents=True, exist_ok=True)
        docx_path = _SOURCE_DOCS_DIR / f"{code_id}.docx"

        if not docx_path.exists():
            self._create_docx_template(code_id, docx_path)
            QMessageBox.information(
                self,
                "تم إنشاء قالب جديد",
                f"تم إنشاء ملف Word للكود  {code_id}\n\n"
                f"المسار:\n{docx_path.resolve()}\n\n"
                "أضف المحتوى الفني واحفظ الملف — سيُدمج تلقائياً في العروض القادمة.",
            )

        try:
            os.startfile(str(docx_path.resolve()))
        except OSError as exc:
            QMessageBox.warning(
                self,
                "تعذّر فتح الملف",
                f"لم يتمكن النظام من فتح الملف:\n{docx_path}\n\n{exc}",
            )

    def _create_docx_template(self, code_id: str, path: Path) -> None:
        """Write a minimal structured Arabic .docx template for code_id."""
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        entry = self._codes.get(code_id, {})
        name_ar = entry.get("activity_name_ar", code_id)
        name_en = entry.get("activity_name_en", "")

        doc = Document()

        # Heading
        h = doc.add_heading(f"{code_id}  —  {name_ar}", level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        if name_en:
            sub = doc.add_paragraph(name_en)
            sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
            sub.runs[0].font.size = Pt(12)
            sub.runs[0].italic = True

        doc.add_paragraph()

        # Standard technical sections
        for section in [
            "الوصف العام",
            "المواصفات الفنية",
            "طريقة التنفيذ",
            "معايير الجودة والفحص",
            "ملاحظات خاصة",
        ]:
            sh = doc.add_heading(section, level=2)
            sh.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            placeholder = doc.add_paragraph("[ أضف المحتوى هنا ]")
            placeholder.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            placeholder.runs[0].font.color.rgb = None  # default color

        doc.save(str(path))

    def _toggle_status(self) -> None:
        current_status = self._status.currentData()
        self._set_combo_data(self._status, "inactive" if current_status == "active" else "active")
        self._sync_toggle_label()

    def _validate_form(self, code_id: str) -> str:
        if not code_id:
            return "معرف الكود مطلوب."
        if not _CODE_ID_RE.match(code_id):
            return "معرف الكود يجب أن يكون مثل: 003-PIP-MAIN."
        if self._is_new and code_id in self._codes:
            return "هذا الكود موجود بالفعل. اختر معرفًا آخر."
        if not self._name_ar.text().strip():
            return "الاسم العربي مطلوب."
        if not self._selected_projects:
            return "اختر مشروعًا واحدًا على الأقل."
        if not self._selected_owners:
            return "اختر جهة مالكة واحدة على الأقل."
        missing_deps = [dep for dep in self._selected_deps if dep not in self._codes]
        if missing_deps:
            return f"توجد تبعيات غير موجودة: {', '.join(missing_deps)}"
        return ""

    def _entry_from_form(self, code_id: str) -> dict[str, Any]:
        today = date.today().isoformat()
        previous = self._codes.get(self._current_code_id or code_id, {})
        return {
            **previous,
            "code_id": code_id,
            "category": self._category.currentData(),
            "phase": self._phase.text().strip().upper() or code_id.split("-")[1],
            "variation": self._variation.text().strip().upper() or code_id.split("-")[2],
            "activity_name_ar": self._name_ar.text().strip(),
            "activity_name_en": self._name_en.text().strip() or self._name_ar.text().strip(),
            "project_ids": self._selected_projects,
            "network_types": self._selected_networks,
            "applicable_owners": self._selected_owners,
            "sequence_order": self._sequence.value(),
            "dependencies": self._selected_deps,
            "page_count": self._pages.value(),
            "source_document": self._source_document.text().strip(),
            "status": self._status.currentData(),
            "last_modified": today,
            "created_date": previous.get("created_date", today),
            "version": previous.get("version", "1.0"),
            "notes": self._notes.toPlainText().strip(),
        }

    def _project_options(self) -> list[tuple[str, str]]:
        projects = self._config_data.get("projects", {})
        return [
            (pid, str(data.get("name_ar") or data.get("project_name_ar") or pid))
            for pid, data in sorted(projects.items())
            if isinstance(data, dict)
        ]

    def _owner_options(self) -> list[tuple[str, str]]:
        owners = self._config_data.get("owner_specifications", {})
        return [
            (oid, str(data.get("owner_name_ar") or data.get("name_ar") or oid))
            for oid, data in sorted(owners.items())
            if isinstance(data, dict)
        ]

    def _network_options(self) -> list[tuple[str, str]]:
        networks = self._config_data.get("network_types", {})
        return [
            (nid, str(data.get("name_ar") or data.get("network_type_name") or nid))
            for nid, data in sorted(networks.items())
            if isinstance(data, dict)
        ]

    def _sync_picker_labels(self) -> None:
        self._projects_btn.setText(f"المشاريع ({len(self._selected_projects)})")
        self._owners_btn.setText(f"الجهات ({len(self._selected_owners)})")
        self._networks_btn.setText(f"الشبكات ({len(self._selected_networks)})")
        self._deps_btn.setText(f"التبعيات ({len(self._selected_deps)})")

    def _sync_toggle_label(self) -> None:
        if self._status.currentData() == "active":
            self._toggle_btn.setText("تعطيل الكود")
        else:
            self._toggle_btn.setText("تفعيل الكود")

    def _select_code(self, code_id: str) -> None:
        for i in range(self._code_list.count()):
            item = self._code_list.item(i)
            if item.data(Qt.UserRole) == code_id:
                self._code_list.setCurrentItem(item)
                break

    @staticmethod
    def _set_combo_data(combo: QComboBox, value: str) -> None:
        idx = combo.findData(value)
        if idx >= 0:
            combo.setCurrentIndex(idx)


def _safe_int(value: Any, default: int) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _string_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item) for item in value if isinstance(item, str)]
    if isinstance(value, str) and value.strip():
        return [value.strip()]
    return []
